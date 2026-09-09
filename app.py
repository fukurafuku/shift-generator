from datetime import date, datetime
import io
import holidays
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from ortools.sat.python import cp_model
import pandas as pd
import streamlit as st

WEEKDAYS_JP = ["月", "火", "水", "木", "金", "土", "日"]


# --- 日本の祝日取得ユーティリティ ---
def get_japanese_holidays(dates):
    if not dates:
        return set()
    years = sorted(set(d.year for d in dates))
    jp_holidays = holidays.JP(years=years)
    return {d for d in dates if d in jp_holidays}


# --- 柔軟な文字判定ユーティリティ ---
def is_maru(val):
    if pd.isna(val):
        return False
    s = str(val).strip()
    return s in [
        "○",
        "〇",
        "o",
        "O",
        "1",
        "True",
        "◯",
        "⚪︎",
        "可",
        "可能",
        "はい",
    ]


def is_batsu(val):
    if pd.isna(val):
        return False
    s = str(val).strip()
    return s in ["×", "x", "X", "不可", "NG", "休", "オフ"]


def is_fulltime_str(val):
    if pd.isna(val):
        return False
    s = str(val).strip()
    return any(k in s for k in ["正社員", "社員", "常勤", "正"])


def is_qualified_str(val):
    if pd.isna(val):
        return False
    s = str(val).strip()
    return any(k in s for k in ["薬剤師", "登録販売者", "資格者", "資格", "有"])


# --- シフト生成エンジン ---
def generate_shift(
    dates_list,
    req_min,
    staff_info,
    holiday_requests,
    extra_work,
    store_closed_days,  # 店舗定休日（曜日リスト）
    is_holiday_open,  # 祝日営業フラグ (True/False)
    jp_holidays_set,  # 祝日日付の集合
    target_off_days=9,
):
    model = cp_model.CpModel()
    staffs = list(staff_info.keys())
    shifts = ["早", "遅", "出", "研", "公休", "有"]

    x = {}
    for s in staffs:
        for d in dates_list:
            d_str = d.strftime("%Y-%m-%d")
            for sh in shifts:
                x[(s, d_str, sh)] = model.NewBoolVar(f"x_{s}_{d_str}_{sh}")

    # 各人各日につき、必ずいずれか1つのシフト
    for s in staffs:
        for d in dates_list:
            d_str = d.strftime("%Y-%m-%d")
            model.AddExactlyOne(x[(s, d_str, sh)] for sh in shifts)

    # 条件設定
    for d in dates_list:
        d_str = d.strftime("%Y-%m-%d")
        w_jp = WEEKDAYS_JP[d.weekday()]
        is_holiday = d in jp_holidays_set

        # 店舗定休日チェック（定休日、または祝日営業でない日の祝日）
        is_store_closed = (w_jp in store_closed_days) or (
            is_holiday and not is_holiday_open
        )

        for s in staffs:
            info = staff_info[s]
            is_holiday_req = s in holiday_requests.get(d_str, [])
            is_extra = s in extra_work.get(d_str, [])
            is_fixed_off = w_jp in info["off_weekdays"]

            # 店舗休業日の場合は全員必ず公休
            if is_store_closed:
                model.Add(x[(s, d_str, "公休")] == 1)
                continue

            # 1. 研修の判定（カレンダーで指定された場合のみ「研」、それ以外は「研」を禁止）
            if is_extra:
                model.Add(x[(s, d_str, "研")] == 1)
            else:
                model.Add(x[(s, d_str, "研")] == 0)

            # 2. 希望休・固定休 ➔ 休日（「公休」または「有」）
            if is_holiday_req or is_fixed_off:
                model.Add(x[(s, d_str, "公休")] + x[(s, d_str, "有")] == 1)

            # 3. 時間帯制約（早番・遅番不可の場合はそれぞれの枠を禁止）
            if not info["can_early"]:
                model.Add(x[(s, d_str, "早")] == 0)
            if not info["can_late"]:
                model.Add(x[(s, d_str, "遅")] == 0)

            # 4. パート（非正社員）の絶対出勤ルール：希望休・固定休でない日は「公休」「有」を禁止
            if not info["is_fulltime"]:
                if not is_holiday_req and not is_fixed_off:
                    model.Add(x[(s, d_str, "公休")] == 0)
                    model.Add(x[(s, d_str, "有")] == 0)

    # 正社員のみ公休数を固定（パートは対象外）
    for s in staffs:
        if staff_info[s]["is_fulltime"]:
            model.Add(
                sum(x[(s, d.strftime("%Y-%m-%d"), "公休")] for d in dates_list)
                == target_off_days
            )

    # 最低必要人数の確保（店舗営業日のみカウント）
    shortage_q_tot, shortage_q_early, shortage_q_late, shortage_unq = (
        {},
        {},
        {},
        {},
    )

    for d in dates_list:
        d_str = d.strftime("%Y-%m-%d")
        w_jp = WEEKDAYS_JP[d.weekday()]
        is_holiday = d in jp_holidays_set
        is_store_closed = (w_jp in store_closed_days) or (
            is_holiday and not is_holiday_open
        )

        q_staffs = [s for s in staffs if staff_info[s]["is_qualified"]]
        unq_staffs = [s for s in staffs if not staff_info[s]["is_qualified"]]

        if is_store_closed:
            # 店舗休業日は必要人数0で固定
            req_qt, req_qe, req_ql, req_u = 0, 0, 0, 0
        else:
            req_qt = req_min["QUALIFIED_TOTAL"].get(w_jp, 0)
            req_qe = req_min["QUALIFIED_EARLY"].get(w_jp, 0)
            req_ql = req_min["QUALIFIED_LATE"].get(w_jp, 0)
            req_u = req_min["UNQUALIFIED"].get(w_jp, 0)

        # 資格者合計
        sqt = model.NewIntVar(0, req_qt, f"sqt_{d_str}")
        model.Add(
            sum(
                x[(s, d_str, sh)]
                for s in q_staffs
                for sh in ["早", "遅", "出"]
            )
            + sqt
            >= req_qt
        )
        shortage_q_tot[d_str] = sqt

        # 資格者早番
        sqe = model.NewIntVar(0, req_qe, f"sqe_{d_str}")
        model.Add(sum(x[(s, d_str, "早")] for s in q_staffs) + sqe >= req_qe)
        shortage_q_early[d_str] = sqe

        # 資格者遅番
        sql = model.NewIntVar(0, req_ql, f"sql_{d_str}")
        model.Add(sum(x[(s, d_str, "遅")] for s in q_staffs) + sql >= req_ql)
        shortage_q_late[d_str] = sql

        # 一般スタッフ
        su = model.NewIntVar(0, req_u, f"su_{d_str}")
        model.Add(
            sum(
                x[(s, d_str, sh)]
                for s in unq_staffs
                for sh in ["早", "遅", "出"]
            )
            + su
            >= req_u
        )
        shortage_unq[d_str] = su

    # 一般スタッフが0人の日に対するペナルティ判定
    zero_unq_days = []
    unq_staffs = [s for s in staffs if not staff_info[s]["is_qualified"]]
    if unq_staffs:
        for d in dates_list:
            d_str = d.strftime("%Y-%m-%d")
            w_jp = WEEKDAYS_JP[d.weekday()]
            is_holiday = d in jp_holidays_set
            is_store_closed = (w_jp in store_closed_days) or (
                is_holiday and not is_holiday_open
            )
            if not is_store_closed:
                is_zero = model.NewBoolVar(f"zero_unq_{d_str}")
                unq_work_count = sum(
                    x[(s, d_str, sh)]
                    for s in unq_staffs
                    for sh in ["早", "遅", "出"]
                )
                model.Add(unq_work_count >= 1).OnlyEnforceIf(is_zero.Not())
                model.Add(unq_work_count == 0).OnlyEnforceIf(is_zero)
                zero_unq_days.append(is_zero)

    # 連続勤務超過ペナルティ
    over_consec = {}
    for s in staffs:
        max_c = staff_info[s]["max_consec"]
        if max_c > 0:
            for i in range(len(dates_list) - max_c):
                window = [
                    dates_list[i + k].strftime("%Y-%m-%d")
                    for k in range(max_c + 1)
                ]
                ov = model.NewIntVar(0, 1, f"ov_{s}_{i}")
                model.Add(
                    sum(
                        x[(s, d_str, sh)]
                        for d_str in window
                        for sh in ["早", "遅", "出", "研"]
                    )
                    <= max_c + ov
                )
                over_consec[(s, window[-1])] = ov

    # 土日両方出勤ペナルティ
    weekend_both_work = []
    for i, d in enumerate(dates_list):
        if d.weekday() == 5:
            if (
                i + 1 < len(dates_list) and dates_list[i + 1].weekday() == 6
            ):  # 日曜日
                d_sat_str = d.strftime("%Y-%m-%d")
                d_sun_str = dates_list[i + 1].strftime("%Y-%m-%d")

                for s in staffs:
                    both_v = model.NewBoolVar(f"both_weekend_{s}_{d_sat_str}")
                    sat_work = sum(
                        x[(s, d_sat_str, sh)] for sh in ["早", "遅", "出", "研"]
                    )
                    sun_work = sum(
                        x[(s, d_sun_str, sh)] for sh in ["早", "遅", "出", "研"]
                    )
                    model.Add(sat_work + sun_work <= 1 + both_v)
                    weekend_both_work.append(both_v)

    total_paid_leaves = sum(
        x[(s, d.strftime("%Y-%m-%d"), "有")]
        for s in staffs
        for d in dates_list
    )

    penalty_shortage = sum(
        shortage_q_tot[d]
        + shortage_q_early[d]
        + shortage_q_late[d]
        + shortage_unq[d]
        for d in shortage_q_tot
    )
    penalty_consec = sum(over_consec.values())
    penalty_weekend = sum(weekend_both_work)
    penalty_zero_unq = sum(zero_unq_days)

    model.Minimize(
        penalty_shortage * 1000
        + penalty_zero_unq * 5
        + penalty_consec * 10
        + penalty_weekend * 5
        + total_paid_leaves
    )

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = 30.0
    status = solver.Solve(model)

    if status in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        grid = {}
        for d in dates_list:
            d_str = d.strftime("%Y-%m-%d")
            w_jp = WEEKDAYS_JP[d.weekday()]
            col_header = f"{d.month}/{d.day}({w_jp})"
            grid[col_header] = {}

            for s in staffs:
                for sh in shifts:
                    if solver.Value(x[(s, d_str, sh)]) == 1:
                        disp_sh = (
                            "休"
                            if sh == "公休"
                            else ("有休" if sh == "有" else sh)
                        )
                        grid[col_header][s] = disp_sh
                        break

            alerts_d = []
            if solver.Value(shortage_q_tot[d_str]) > 0:
                alerts_d.append(
                    f"資格者計不足-{solver.Value(shortage_q_tot[d_str])}"
                )
            if solver.Value(shortage_q_early[d_str]) > 0:
                alerts_d.append(
                    f"資格早番不足-{solver.Value(shortage_q_early[d_str])}"
                )
            if solver.Value(shortage_q_late[d_str]) > 0:
                alerts_d.append(
                    f"資格遅番不足-{solver.Value(shortage_q_late[d_str])}"
                )
            if solver.Value(shortage_unq[d_str]) > 0:
                alerts_d.append(
                    f"一般不足-{solver.Value(shortage_unq[d_str])}"
                )

            grid[col_header]["【日別人数不足】"] = (
                " / ".join(alerts_d) if alerts_d else "正常"
            )

        grid["【個人別集計】"] = {}
        for s in staffs:
            off_c = sum(
                1
                for d in dates_list
                if solver.Value(x[(s, d.strftime("%Y-%m-%d"), "公休")]) == 1
            )
            paid_c = sum(
                1
                for d in dates_list
                if solver.Value(x[(s, d.strftime("%Y-%m-%d"), "有")]) == 1
            )
            summary_str = f"公休:{off_c}日"
            if paid_c > 0:
                summary_str += f" / 有休:{paid_c}日"

            grid["【個人別集計】"][s] = summary_str

        grid["【個人別集計】"]["【日別人数不足】"] = "-"
        return pd.DataFrame(grid)
    return None


# --- 装飾付きExcelバイナリ生成関数（Streamlitダウンロード用） ---
def get_colored_excel_bytes(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "シフト表"

    # 色の定義
    fill_early = PatternFill(
        start_color="FFE6CC", end_color="FFE6CC", fill_type="solid"
    )  # 薄いオレンジ（早）
    fill_late = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )  # 薄い紺/ソフトブルー（遅）
    fill_work = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )  # 薄い緑（出）

    # ヘッダー書き込み
    headers = ["スタッフ / 日付"] + list(df.columns)
    ws.append(headers)

    # データ書き込み＆色付け
    for index, row in df.iterrows():
        row_data = [index] + list(row.values)
        ws.append(row_data)

    for r_idx, row in enumerate(
        ws.iter_rows(
            min_row=2,
            max_row=len(df) + 1,
            min_col=2,
            max_col=len(df.columns) + 1,
        ),
        start=0,
    ):
        for c_idx, cell in enumerate(row):
            val = str(cell.value).strip() if cell.value is not None else ""
            if val == "早":
                cell.fill = fill_early
            elif val == "遅":
                cell.fill = fill_late
            elif val == "出":
                cell.fill = fill_work

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# --- Streamlit メインアプリ ---
st.title("📂 シフト自動作成アプリ")

uploaded_file = st.file_uploader(
    "シフト作成Excelファイルをアップロードしてください", type=["xlsx"]
)

if uploaded_file is not None:
    excel_file = io.BytesIO(uploaded_file.read())

    # 1. 固定ルールシートの読み取り
    df_fix = pd.read_excel(excel_file, sheet_name="固定ルール", header=None)

    col_map_days = {"月": 2, "火": 3, "水": 4, "木": 5, "金": 6, "土": 7, "日": 8}
    store_closed_days = []
    is_holiday_open = True  # デフォルト営業

    if len(df_fix) > 5:
        for w_key, col_i in col_map_days.items():
            if col_i < len(df_fix.columns):
                val = df_fix.iloc[5, col_i]
                if is_batsu(val):
                    store_closed_days.append(w_key)

        if len(df_fix.columns) > 9:
            val_j6 = df_fix.iloc[5, 9]
            if is_batsu(val_j6):
                is_holiday_open = False

    # 必要人数の取得
    req_min = {
        "QUALIFIED_TOTAL": {w: 0 for w in WEEKDAYS_JP},
        "QUALIFIED_EARLY": {w: 0 for w in WEEKDAYS_JP},
        "QUALIFIED_LATE": {w: 0 for w in WEEKDAYS_JP},
        "UNQUALIFIED": {w: 0 for w in WEEKDAYS_JP},
    }

    for r in range(len(df_fix)):
        label = (
            str(df_fix.iloc[r, 1]).strip()
            if pd.notna(df_fix.iloc[r, 1])
            else ""
        )
        if "資格者合計" in label:
            for w_key, col_i in col_map_days.items():
                if col_i < len(df_fix.columns) and pd.notna(df_fix.iloc[r, col_i]):
                    try:
                        req_min["QUALIFIED_TOTAL"][w_key] = int(
                            df_fix.iloc[r, col_i]
                        )
                    except ValueError:
                        pass
        elif "うち資格者早番" in label:
            for w_key, col_i in col_map_days.items():
                if col_i < len(df_fix.columns) and pd.notna(df_fix.iloc[r, col_i]):
                    try:
                        req_min["QUALIFIED_EARLY"][w_key] = int(
                            df_fix.iloc[r, col_i]
                        )
                    except ValueError:
                        pass
        elif "うち資格者遅番" in label:
            for w_key, col_i in col_map_days.items():
                if col_i < len(df_fix.columns) and pd.notna(df_fix.iloc[r, col_i]):
                    try:
                        req_min["QUALIFIED_LATE"][w_key] = int(
                            df_fix.iloc[r, col_i]
                        )
                    except ValueError:
                        pass
        elif "一般スタッフ" in label:
            for w_key, col_i in col_map_days.items():
                if col_i < len(df_fix.columns) and pd.notna(df_fix.iloc[r, col_i]):
                    try:
                        req_min["UNQUALIFIED"][w_key] = int(
                            df_fix.iloc[r, col_i]
                        )
                    except ValueError:
                        pass

    # スタッフ一覧の読み込み
    header_row_idx = 15
    for idx, row in df_fix.iterrows():
        row_vals = [str(v).strip() for v in row.values if pd.notna(v)]
        if any(k in row_vals for k in ["スタッフ一覧", "スタッフ", "名前", "氏名"]):
            header_row_idx = idx
            break

    excel_file.seek(0)
    df_staff = pd.read_excel(
        excel_file, sheet_name="固定ルール", skiprows=header_row_idx
    )
    staff_info = {}
    maru_symbols = ["①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨"]
    w_indices = {
        "月": 6,
        "火": 7,
        "水": 8,
        "木": 9,
        "金": 10,
        "土": 11,
        "日": 12,
    }

    idx_count = 0
    for _, row in df_staff.iterrows():
        s_val = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
        if not s_val or s_val in ["nan", "None", "", "スタッフ名", "名前", "氏名"]:
            continue

        s_id = (
            maru_symbols[idx_count]
            if idx_count < len(maru_symbols)
            else f"スタッフ{idx_count+1}"
        )

        type_val = (
            str(row.iloc[2]).strip()
            if len(row) > 2 and pd.notna(row.iloc[2])
            else ""
        )
        job_val = (
            str(row.iloc[3]).strip()
            if len(row) > 3 and pd.notna(row.iloc[3])
            else ""
        )

        can_e = is_maru(row.iloc[4]) if len(row) > 4 else False
        can_l = is_maru(row.iloc[5]) if len(row) > 5 else False
        is_ft = is_fulltime_str(type_val)
        is_q = is_qualified_str(job_val)

        off_weekdays = [
            w
            for w, c in w_indices.items()
            if c < len(row) and is_batsu(row.iloc[c])
        ]
        try:
            max_consec = (
                int(row.iloc[13])
                if len(row) > 13 and pd.notna(row.iloc[13])
                else 6
            )
        except ValueError:
            max_consec = 6

        staff_info[s_id] = {
            "name": s_val,
            "is_fulltime": is_ft,
            "is_qualified": is_q,
            "can_early": can_e,
            "can_late": can_l,
            "off_weekdays": off_weekdays,
            "max_consec": max_consec,
        }
        idx_count += 1

    # 2. カレンダー入力シートの読み取り
    excel_file.seek(0)
    df_cal = pd.read_excel(excel_file, sheet_name="カレンダー入力", header=None)

    # H2セルから設定公休数の取得
    target_off_days = 9
    if len(df_cal) > 1 and len(df_cal.columns) > 7:
        val_h2 = df_cal.iloc[1, 7]
        if pd.notna(val_h2):
            try:
                target_off_days = int(val_h2)
            except ValueError:
                pass

    dynamic_staff_cols = {}
    cal_header_row = 4
    for r in range(2, 6):
        if r < len(df_cal):
            row_vals = [
                str(v).strip() for v in df_cal.iloc[r].values if pd.notna(v)
            ]
            if any(s_id in row_vals for s_id in staff_info.keys()):
                cal_header_row = r
                break

    for c in range(len(df_cal.columns)):
        val = (
            str(df_cal.iloc[cal_header_row, c]).strip()
            if pd.notna(df_cal.iloc[cal_header_row, c])
            else ""
        )
        if val in staff_info:
            dynamic_staff_cols[val] = c

    holiday_requests, extra_work, dates_list = {}, {}, []

    for r in range(cal_header_row + 1, len(df_cal)):
        d_val = df_cal.iloc[r, 3]
        if pd.isna(d_val) or str(d_val).strip() in ["", "nan", "None"]:
            continue
        try:
            if isinstance(d_val, (datetime, date)):
                d_obj = d_val if isinstance(d_val, date) else d_val.date()
            else:
                dt = pd.to_datetime(str(d_val).strip())
                d_obj = dt.date()
        except Exception:
            continue

        d_str = d_obj.strftime("%Y-%m-%d")
        dates_list.append(d_obj)
        holiday_requests[d_str] = []
        extra_work[d_str] = []

        for s_id, col_idx in dynamic_staff_cols.items():
            if col_idx < len(df_cal.columns):
                cell_val = (
                    str(df_cal.iloc[r, col_idx]).strip()
                    if pd.notna(df_cal.iloc[r, col_idx])
                    else ""
                )
                if cell_val in ["公休", "休", "希望休", "×"]:
                    holiday_requests[d_str].append(s_id)
                elif cell_val in ["研", "研修"]:
                    extra_work[d_str].append(s_id)

    dates_list = sorted(list(set(dates_list)))
    jp_holidays_set = get_japanese_holidays(dates_list)

    if dates_list:
        st.info(
            f"📌 カレンダー入力H2から取得した設定公休数: {target_off_days}日 | "
            f"店舗定休日: {store_closed_days if store_closed_days else 'なし'} | "
            f"祝日営業設定: {'営業' if is_holiday_open else '休業'}"
        )

        with st.spinner("⚙️ シフト自動生成を実行中..."):
            result_df = generate_shift(
                dates_list,
                req_min,
                staff_info,
                holiday_requests,
                extra_work,
                store_closed_days,
                is_holiday_open,
                jp_holidays_set,
                target_off_days,
            )

        if result_df is not None:
            st.success("🎉 シフト表の作成が完了しました！")
            st.dataframe(result_df)

            out_name = f"完成シフト表_{dates_list[0]}_{dates_list[-1]}.xlsx"
            excel_data = get_colored_excel_bytes(result_df)

            st.download_button(
                label="📥 完成したExcelファイルをダウンロード",
                data=excel_data,
                file_name=out_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        else:
            st.error("条件を満たすシフトが見つかりませんでした。")
